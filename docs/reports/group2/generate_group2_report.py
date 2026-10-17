#!/usr/bin/env python3
"""Build the Group 2 Academic & Course report as DOCX and PDF.

All project statements are sourced from the repository. UI figures are actual React
screens captured with the explicit synthetic API fixtures in capture_screenshots.cjs.
"""
from pathlib import Path
from html import escape as html_escape
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    BaseDocTemplate, Frame, PageTemplate, Paragraph, Spacer, PageBreak,
    Table, TableStyle, Image as RLImage, KeepTogether, NextPageTemplate,
)
from reportlab.pdfgen.canvas import Canvas
from reportlab.platypus.tableofcontents import TableOfContents
from reportlab.lib.utils import ImageReader

ROOT = Path(__file__).resolve().parents[3]
REPORT_DIR = Path(__file__).resolve().parent
SCREEN_DIR = REPORT_DIR / 'screenshots'
LOGO = ROOT / 'frontend/public/brand/pict-logo-256.png'
ER_IMAGE = REPORT_DIR / 'ER_Module_2_Core.png'
DOCX_OUT = ROOT / 'docs/reports/PICT_ERP_Group_02_Academic_Course_Report.docx'
PDF_OUT = ROOT / 'docs/reports/PICT_ERP_Group_02_Academic_Course_Report.pdf'

NAVY = '16365C'
BLUE = '2D5F91'
TEAL = '0F766E'
ORANGE = 'B45309'
PALE = 'EEF4FA'
PALE2 = 'F4F7FB'
INK = '243447'
MUTED = '5B6B7A'
BORDER = 'D4DDE7'


def B(kind, **kwargs):
    return {'type': kind, **kwargs}


CONTENT = [
    B('h1', text='1. Executive Summary'),
    B('p', text=(
        'The University Higher Education ERP System is a database-centric, seven-module project built around a shared relational schema. '
        'Group 2 (Academic & Course) provides the academic reference layer: departments, programs, batches, sections, subjects, curriculum mapping, '
        'course offerings, enrolments, faculty-subject allocations, timetabling and the academic calendar. The module is important because other '
        'areas attach their work to the course offering: attendance sessions, examination registrations, enrolments and faculty teaching assignments.'
    )),
    B('p', text=(
        'The repository implements the module through React 18 and TypeScript screens, an Express/TypeScript API, and MySQL 8 / MariaDB-compatible '
        'SQL. It puts enrolment capacity, duplicate prevention and timetable validation in database constraints, a transaction procedure and triggers, '
        'rather than trusting only the browser. This report describes the repository as inspected and separates documented project claims from evidence '
        'actually available during preparation.'
    )),
    B('callout', text=(
        'Evidence note: The six interface figures in this report are authentic browser captures of the repository React application. The API and database '
        'were unavailable in the report sandbox, so Playwright intercepted /api/v1 requests and supplied explicitly synthetic illustrative responses. '
        'These are not live-data screenshots and do not verify backend or database behaviour.'
    )),
    B('table', headers=['At a glance', 'Repository-backed description'], rows=[
        ['Group 2 table inventory', '13 academic and course tables listed in PROJECT_REPORT.md §21 and MODULE_INTEGRATION.md §3'],
        ['Primary UI pages', 'Academic masters, Subjects, Course offerings, Timetable and Academic calendar'],
        ['Key DBMS logic', 'sp_bulk_enroll_students; enrolment guard/counter triggers; timetable conflict trigger'],
        ['Repository scale (all modules)', 'Documentation reports 61 tables, 20 views, 17 procedures, 13 functions, 28 triggers and 138 foreign keys'],
        ['Current year shown in fixtures', 'A.Y. 2026–27, Semester 5, matching the repository seed flags; not freshly queried from a database'],
    ]),

    B('h1', text='2. Project Context, Problem and Objectives'),
    B('h2', text='2.1 Problem statement'),
    B('p', text=(
        'Academic data is often spread across department spreadsheets and separate systems. That makes it difficult to keep curriculum definitions, '
        'semester offerings, section membership, faculty assignments and room schedules consistent. In a joined ERP, one course-offering record can be '
        'referenced by enrolment, attendance, examinations and timetable functions, reducing duplicated subject or faculty data.'
    )),
    B('h2', text='2.2 Group 2 objectives'),
    B('bullets', items=[
        'Maintain a single department/program/batch/section structure for the rest of the ERP to reference.',
        'Maintain subject metadata and map subjects into each program and semester through program_subjects.',
        'Create semester- and section-specific course offerings that connect subjects, programs, faculty and capacity.',
        'Enrol students individually or in bulk while enforcing duplicate and capacity rules.',
        'Schedule classes by day, start/end time and room, and reject defined faculty, room and section clashes.',
        'Publish term dates, holidays, examinations and institute events through academic_calendar.',
        'Provide authenticated, role-aware pages and administrator-controlled write operations.',
    ]),
    B('h2', text='2.3 Scope boundaries'),
    B('table', headers=['In scope for Group 2', 'Related but owned elsewhere'], rows=[
        ['Departments, programs, batches, sections, years, semesters, subjects, program-subject mappings, course offerings, enrolments, faculty-subject mapping, timetable, academic calendar.',
         'Student identity/admission, faculty HR/payroll, attendance marking, examination schedules/marks/results, fee collection and hostel records. These consume shared academic keys but are not Group 2 CRUD domains.'],
    ]),

    B('h1', text='3. Architecture and Technology'),
    B('p', text=(
        'The repository uses a layered browser-to-database path. The React client calls relative /api/v1 URLs; Vite proxies /api to the Express service in development. '
        'Express routes apply authentication, role checks and Zod request validation, then call controllers, services and parameterised-SQL repositories. '
        'The backend uses mysql2/promise to reach the university_erp schema. The project documents MySQL 8 / MariaDB 10.6+ compatibility; a MySQL/MariaDB service was not available in the report-building sandbox.'
    )),
    B('table', headers=['Layer', 'Technology / role'], rows=[
        ['Frontend', 'React 18 · TypeScript · Vite · Tailwind CSS · React Router · Lucide icons · Recharts'],
        ['Development proxy', 'Vite /api proxy; browser code uses relative API URLs'],
        ['Backend', 'Node 20 · Express · TypeScript (strict) · Zod validation · mysql2/promise'],
        ['Authentication and access', 'JWT authentication; role middleware on protected routes; administrator-only writes for Group 2'],
        ['Database', 'MySQL 8 / MariaDB 10.6+ · InnoDB · FK/unique/check constraints · transactions · procedures · triggers · views'],
        ['Documentation artifacts', 'Generated Markdown data dictionary, Mermaid source and SVG ER diagrams'],
    ]),
    B('h2', text='3.1 Request path'),
    B('code', text='Browser (React) → relative /api/v1 request → Vite /api proxy → Express route\n→ authenticate / role / Zod validation → controller → academic service\n→ repository with bound SQL parameters → MySQL 8 / MariaDB'),
    B('p', text=(
        'This separation keeps business rules out of React and keeps SQL in repository/data-access code. The bulk-enrolment service calls the stored procedure; '
        'the timetable service delegates conflict refusal to the database trigger. UI permissions improve navigation, while the API role checks are the security boundary.'
    )),

    B('h1', text='4. Requirements and User Roles'),
    B('h2', text='4.1 Functional requirements'),
    B('table', headers=['ID', 'Requirement', 'Primary implementation'], rows=[
        ['FR-01', 'Maintain departments, programs, batches, sections, years and semesters.', 'AcademicsMastersPage; /academics/* reference endpoints; relational master tables.'],
        ['FR-02', 'Maintain subject catalog with credits, type and contact hours.', 'SubjectsPage; subjects table; subject-code uniqueness.'],
        ['FR-03', 'Map subjects to programs/semesters and define delivered offerings.', 'program_subjects schema bridge and course_offerings API. No curriculum route was found in the inspected academic.routes.ts; a route/UI curriculum workflow is not verified.'],
        ['FR-04', 'Enroll a student in an offering and bulk-enrol a program/semester cohort.', 'POST /enrollments; POST /enrollments/bulk; sp_bulk_enroll_students.'],
        ['FR-05', 'Schedule course slots and expose conflict information.', 'TimetablePage; timetable table; trigger and GET /timetable/conflicts.'],
        ['FR-06', 'Maintain events and holidays by academic year and optional semester.', 'CalendarPage; academic_calendar table and endpoints.'],
        ['FR-07', 'Audit administrator create/update/delete actions.', 'Academic service audit calls; audit_logs table (as documented).'],
        ['FR-08', 'Expose shared offerings to attendance, examination and faculty workflows.', 'FK-based integration through course_offerings and shared student/faculty masters.'],
    ]),
    B('h2', text='4.2 User access model'),
    B('table', headers=['Actor / role', 'Group 2 use (as wired in routes)'], rows=[
        ['ADMIN', 'May create/update academic subjects and offerings, create sections, set class teacher, enrol/drop, assign faculty-subjects, add/remove timetable slots and calendar events.'],
        ['FACULTY', 'Authenticated reads for academic reference data; faculty timetable and faculty-subject reads are role-scoped.'],
        ['EXAM_CELL, ACCOUNTANT, HOSTEL_ADMIN, HR', 'Authenticated read access to selected academic reference lists; some routes explicitly permit these staff roles.'],
        ['STUDENT', 'Student-specific timetable endpoint is allowed for the student role; general reference routes require authentication. No Group 2 write route is exposed.'],
    ]),
    B('p', text=(
        'The route file explicitly attaches requireRole(\'ADMIN\') to Group 2 write endpoints. Read access is not identical across endpoints: some require only authentication, '
        'some use staff-role lists, and the student timetable route has its own allowed-role list. Verify each route before documenting a broader role guarantee.'
    )),

    B('h1', text='5. Group 2 Module Design and Workflows'),
    B('h2', text='5.1 Academic masters and curriculum'),
    B('p', text=(
        'Departments own programs and subject catalog entries. Batches identify cohorts and their academic-year intake; sections group students and can refer to a class teacher. '
        'program_subjects is the intersection table for a program-to-subject many-to-many relationship, with semester_no and is_elective attributes. Its composite unique key '
        'prevents repeating the same subject in the same program semester. The repository also contains curriculum service methods, but the inspected academic.routes.ts does not mount a curriculum route; this report treats the table as present and the end-to-end curriculum CRUD path as unverified.'
    )),
    B('h2', text='5.2 Course offering as the delivery pivot'),
    B('p', text=(
        'A course_offerings row binds a subject to a semester and academic year, program, optional section and optional faculty member, and stores capacity, enrolled_count, offering_type and status. '
        'The offering is the operational key used downstream by enrolments, attendance sessions, examination registration and timetable slots. The row therefore represents a course actually delivered '
        'to a cohort, not just a curriculum catalog item.'
    )),
    B('h2', text='5.3 Enrolment workflow'),
    B('numbered', items=[
        'An administrator selects a program, semester and optional section for bulk processing.',
        'The API requires authentication and the ADMIN role, then calls the academic service.',
        'The service invokes sp_bulk_enroll_students with the program, semester, section and actor.',
        'The procedure starts a transaction, counts existing pairs, checks offering capacity and stages new student/offering pairs in a temporary table.',
        'It inserts eligible missing pairs, returns enrolled/skipped/message outputs, writes an audit entry and reconciles offering counters.',
        'The unique student/offering key and enrolment guard protect against duplicates, cancelled offerings and capacity violations at the database layer.',
    ]),
    B('h2', text='5.4 Timetable and conflict handling'),
    B('p', text=(
        'A timetable row identifies an offering, section and faculty member and records day_of_week, start_time, end_time and room_number. The BEFORE INSERT trigger refuses an end time not later than the start time and '
        'signals readable conflicts when the same faculty member, room or section already has a row with the same day and start time. Unique keys also cover each faculty/day/start, room/day/start and section/day/start combination. '
        'The inspected trigger checks identical start-time slots; it should not be described as a general interval-overlap detector.'
    )),
    B('h2', text='5.5 Academic calendar'),
    B('p', text=(
        'The academic calendar holds title, description, event_type, start/end dates, holiday flag and a required academic_year_id with an optional semester_id. The UI displays upcoming events first and lists the full event record below.'
    )),

    B('h1', text='6. Data Model and ER Diagram'),
    B('image', path=ER_IMAGE, caption='Figure 1. Repository-generated Group 2 core ER diagram. The report shortens the subtitle for print; the original SVG/Mermaid sources remain under docs/er/.', width=6.70),
    B('callout', text=(
        'Diagram coverage note: docs/ER_DIAGRAMS.md records 12 entities and 14 relationships for the generated Module 2 SVG. The current Group 2 table inventory in PROJECT_REPORT.md §21 and MODULE_INTEGRATION.md §3 has 13 module tables. The generated picture includes shared students and faculty but omits batches, faculty_subjects and academic_calendar. This report retains the repository diagram and lists the omitted tables and their keys below; treat the graphic as the core diagram, not the complete physical schema.'
    )),
    B('h2', text='6.1 Group 2 table inventory'),
    B('table', headers=['Table', 'Role in the data model', 'Primary and selected foreign keys'], rows=[
        ['departments', 'Shared organizational master.', 'PK department_id; referenced by programs, subjects and faculty.'],
        ['programs', 'Academic program definition.', 'PK program_id; FK department_id; unique program_code.'],
        ['batches', 'Cohort / intake year.', 'PK batch_id; FK program_id and academic_year_id; unique batch_code.'],
        ['sections', 'Teaching group within a program/batch.', 'PK section_id; FK program_id, batch_id, optional semester_id and class_teacher_id.'],
        ['subjects', 'Catalog of taught subject units.', 'PK subject_id; FK department_id; unique subject_code.'],
        ['program_subjects', 'Curriculum mapping bridge.', 'PK program_subject_id; FKs program_id, subject_id; unique program/subject/semester.'],
        ['course_offerings', 'Subject delivered to a term/program/section.', 'PK offering_id; FKs subject, semester, year, program, optional section/faculty.'],
        ['enrollments', 'Student registration against an offering.', 'PK enrollment_id; FKs student_id and offering_id; unique student/offering.'],
        ['faculty_subjects', 'Term-bounded teaching allocation.', 'PK allocation_id; FKs faculty, subject, academic_year and optional semester.'],
        ['timetable', 'Recurring weekly slot.', 'PK timetable_id; FKs offering, section and faculty.'],
        ['academic_calendar', 'Academic-year event/holiday.', 'PK event_id; FK academic_year_id and optional semester_id.'],
        ['academic_years', 'Year master/current-year flag.', 'PK academic_year_id; unique year_label.'],
        ['semesters', 'Term within academic year.', 'PK semester_id; FK academic_year_id; unique year/semester_no.'],
    ]),
    B('h2', text='6.2 Core relationship map'),
    B('table', headers=['Parent / master', 'Child / relationship', 'Meaning'], rows=[
        ['departments', 'programs; subjects; faculty', 'One department may own multiple programs, subjects and faculty records.'],
        ['programs', 'batches; sections; program_subjects; course_offerings', 'Program defines the cohort and curriculum/teaching context.'],
        ['academic_years', 'semesters; batches; course_offerings; faculty_subjects; academic_calendar', 'Year is the temporal parent for term and year-specific records.'],
        ['semesters', 'sections; course_offerings; faculty_subjects; academic_calendar', 'Term context; nullable where the schema allows it.'],
        ['subjects', 'program_subjects; course_offerings; faculty_subjects', 'Catalog subjects map into a curriculum and are assigned/offered.'],
        ['batches', 'sections; students', 'A cohort groups students and its sections.'],
        ['sections', 'course_offerings; timetable', 'Sections receive offerings and weekly slots.'],
        ['students', 'enrollments', 'Student/offering many-to-many association is stored in enrollments.'],
        ['faculty', 'course_offerings; faculty_subjects; sections; timetable', 'Faculty may teach offerings, hold subject allocations, serve as class teacher and appear on slots.'],
        ['course_offerings', 'enrollments; timetable', 'Operational course record is the pivot for enrolment and schedule slots.'],
    ]),

    B('pagebreak'),
    B('h1', text='7. Data Dictionary, Keys and Integrity'),
    B('p', text=(
        'The project data dictionary is generated from the schema and documents key types, nullability, indexes and foreign keys. The excerpts below summarize selected Group 2 constraints; they do not replace the full dictionary or a live SHOW CREATE TABLE review.'
    )),
    B('table', headers=['Integrity rule', 'Schema mechanism', 'Effect'], rows=[
        ['Academic-year identity', 'UNIQUE academic_years.year_label', 'Prevents duplicate year labels.'],
        ['Semester within year', 'FK semesters.academic_year_id; UNIQUE (academic_year_id, semester_no)', 'Prevents duplicate semester numbers inside a year.'],
        ['Program and subject masters', 'Unique program_code / subject_code; department FKs', 'Avoids duplicate catalog codes and orphan departments.'],
        ['Curriculum mapping', 'UNIQUE (program_id, subject_id, semester_no)', 'Prevents repeating a subject in a program term.'],
        ['Section identity', 'UNIQUE (batch_id, section_code)', 'Prevents duplicate section labels inside a batch.'],
        ['Course offering', 'Unique subject/semester/program/section tuple; FKs to academic masters', 'Prevents duplicate offering definitions for the same key tuple.'],
        ['Enrolment identity', 'UNIQUE (student_id, offering_id)', 'A student cannot have two enrolment rows for the same offering.'],
        ['Capacity counter', 'enrolled_count plus AFTER INSERT/DELETE counter triggers', 'Keeps the stored offering seat count synchronized with enrolment rows.'],
        ['Faculty subject allocation', 'UNIQUE (faculty_id, subject_id, academic_year_id, semester_id)', 'Prevents duplicate allocation for that faculty, subject and term.'],
        ['Timetable slot', 'Unique faculty/day/start, room/day/start and section/day/start keys', 'Prevents duplicate start slots for the same resource.'],
        ['Calendar integrity', 'FK to academic year and optional semester; ENUM event_type', 'Restricts event associations to existing academic time masters.'],
    ]),
    B('h2', text='7.1 Relational design notes'),
    B('bullets', items=[
        'program_subjects resolves the many-to-many relationship between programs and subjects while keeping semester number and elective status on the mapping.',
        'enrollments resolves the many-to-many relationship between students and course offerings and carries enrollment date/status/grade fields.',
        'course_offerings.enrolled_count is a denormalized counter documented as maintained by insert/delete triggers and the bulk procedure; source rows remain enrollments.',
        'Foreign keys and unique keys are declared in the schema/constraint scripts and summarized in docs/DATA_DICTIONARY.md.',
        'The data dictionary describes lifecycle status columns, defaults, timestamp fields and database indexes; refer to the generated dictionary for exact column types.',
    ]),

    B('h1', text='8. Database Procedures, Triggers and Transactions'),
    B('table', headers=['DBMS object', 'Observed responsibility in source'], rows=[
        ['sp_bulk_enroll_students', 'Starts a transaction; checks existing pairs and capacity; stages missing pairs; inserts enrolments; returns enrolled/skipped/message; handles SQL exceptions with rollback/resignal.'],
        ['trg_enrollments_bi_guard', 'Before insert: reads offering state/capacity and student state; refuses missing/cancelled offering, disallowed student state and full offering.'],
        ['trg_enrollments_ai_count', 'After insert: recomputes the offering enrolled_count from enrollments.'],
        ['trg_enrollments_ad_count', 'After delete: recomputes the offering enrolled_count from enrollments.'],
        ['trg_timetable_bi_conflict', 'Before insert: rejects end_time not after start_time, and same day/start-time faculty, room or section clashes.'],
        ['Unique indexes', 'Database uniqueness provides a second layer for duplicate enrolment and the three timetable resource/start combinations.'],
    ]),
    B('h2', text='8.1 Why these rules belong in the database'),
    B('p', text=(
        'A database guard applies to API calls and direct SQL writes, not just the React forms. The bulk operation has an explicit transaction and error handler, while the timetable trigger raises readable SQLSTATE 45000 errors. '
        'The API service converts repository/procedure outcomes into the shared response envelope and audit messages (as documented in the service code).'
    )),
    B('h2', text='8.2 Group 2 query/report opportunities'),
    B('bullets', items=[
        'Offerings by program, semester, subject, faculty and section.',
        'Seat utilization: enrolled_count divided by capacity, with attention to zero capacity.',
        'Curriculum completion matrix: program_subjects compared with active offerings.',
        'Faculty teaching load by year/semester from faculty_subjects and course_offerings.',
        'Weekly room and section schedules from timetable.',
        'Calendar events grouped by academic year, semester, type and holiday flag.',
    ]),

    B('h1', text='9. API Surface and Security'),
    B('table', headers=['Route group', 'Methods / examples', 'Access noted in route file'], rows=[
        ['/academics/departments, /programs, /batches, /years, /semesters, /sections', 'GET lists/details; current semester lookup.', 'Authenticated reads.'],
        ['/academics/subjects', 'GET list/detail; POST create; PUT update.', 'Authenticated reads; ADMIN writes.'],
        ['/academics/offerings', 'GET list/detail; POST create; PUT update.', 'Authenticated reads; ADMIN writes.'],
        ['/academics/enrollments', 'GET list; POST single; POST /bulk; DELETE by id.', 'GET restricted to listed staff roles; ADMIN writes.'],
        ['/academics/faculty-subjects', 'GET list; POST assign; DELETE allocation.', 'GET ADMIN/FACULTY; writes ADMIN.'],
        ['/academics/timetable', 'GET section/faculty/student; GET conflicts; POST slot; DELETE slot.', 'Student timetable route permits ADMIN/FACULTY/STUDENT; conflicts ADMIN; slot writes ADMIN.'],
        ['/academics/calendar', 'GET list; POST event; DELETE event.', 'Authenticated read; ADMIN writes.'],
    ]),
    B('p', text=(
        'Route source: backend/src/routes/academic.routes.ts. The base prefix is /api/v1/academics because the router is mounted at /academics. API write routes use authenticate and requireRole(\'ADMIN\'); controller validation and service logic are additional layers. The browser uses relative /api/v1 calls, which is appropriate for the Vite proxy and preview host.'
    )),
    B('h2', text='9.1 Security controls described by the repository'),
    B('bullets', items=[
        'JWT authentication and role middleware on protected endpoints.',
        'Zod schemas on write endpoints to validate request bodies before service logic.',
        'Prepared/bound SQL parameters in repository methods rather than concatenated user values.',
        'Administrator-only API writes for the Group 2 setup and scheduling operations.',
        'Audit calls on major academic create/update/delete actions, as implemented in academic.service.ts.',
        'The in-app route/UI restriction is not a substitute for the server-side checks; API authorization is the actual boundary.',
    ]),

    B('h1', text='10. User Interface Evidence'),
    B('p', text=(
        'The following images show the real React application and shared ERP shell, including the available official PICT logo asset. The capture script injects a synthetic ADMIN session token only for local rendering and intercepts API requests with fixed illustrative fixtures. No API mutation was submitted and no database query ran.'
    )),
    B('image_page', path=SCREEN_DIR / '01_Admin_Dashboard_UI.png', caption='Figure 2. Admin dashboard with illustrative synthetic KPI/chart/activity responses. The 2026–27 / Semester 5 label follows the repository seed; all displayed counts and activities are sample fixture values, not a live dashboard query. The chart’s ECE code denotes the separate Electronics and Computer Engineering department (short name E&CE); ETC is the E&TC department code.'),
    B('image_page', path=SCREEN_DIR / '02_Academic_Masters_UI.png', caption='Figure 3. Academic masters page showing the shared department, program, batch and section structure. Email values use example.com demo placeholders; no genuine student/faculty identity is shown.'),
    B('image_page', path=SCREEN_DIR / '03_Subjects_UI.png', caption='Figure 4. Subjects page with an illustrative subset of seeded subject codes and names; the page displays eight rows in this capture, not a fresh query of the 48-subject seed.'),
    B('image_page', path=SCREEN_DIR / '04_Course_Offerings_UI.png', caption='Figure 5. Course offerings page with five synthetic rows. Pagination metadata shows the repository-documented 345-offering seed scale, but the visible records and counts were supplied by the synthetic fixture.'),
    B('image_page', path=SCREEN_DIR / '05_Timetable_UI.png', caption='Figure 6. Weekly timetable cards with “Demo Faculty” labels and sample rooms/slots. No faculty member or live timetable record was queried.'),
    B('image_page', path=SCREEN_DIR / '06_Academic_Calendar_UI.png', caption='Figure 7. Academic calendar page. Event titles/dates reflect repository seed examples for illustration, but the rendered API response was intercepted and is not a live calendar query.'),
    B('pagebreak'),

    B('h1', text='11. Cross-Module Integration'),
    B('table', headers=['Partner module', 'Group 2 integration point', 'Boundary'], rows=[
        ['Student & Admission', 'Admission assigns department/program/batch/section; enrolments reference students.', 'Student identity and approval belong to Module 1.'],
        ['Faculty', 'Offerings and timetable reference faculty_id; faculty_subjects records subject allocations.', 'Faculty HR records and payroll belong to Module 7.'],
        ['Attendance', 'Class sessions attach to course_offerings; attendance guard expects student enrolment for the offering.', 'Session creation/attendance marking belong to Module 3.'],
        ['Exams & Results', 'Exam registrations reference course_offerings; the attendance eligibility logic reads offering enrolment/attendance.', 'Exam schedules, marks and results belong to Module 4.'],
        ['Fees & Finance', 'Fee structures are associated with program and semester in the project integration documentation.', 'Billing/payment processing belongs to Module 5; do not infer a verified exam fee gate from the documentation alone.'],
        ['Reports / Dashboard', 'Group 2 tables feed the current-term header, offering list, calendar and cross-module reporting.', 'The screenshots in this report use fixtures rather than live read models.'],
    ]),
    B('callout', text=(
        'Cross-module verification issue: MODULE_INTEGRATION.md and README.md state that outstanding dues block exam registration. In the inspected database/05_procedures.sql implementation, sp_register_student_for_exam calculates v_dues, but the visible eligibility branch checks attendance and the insert path does not test v_dues > 0. The inspected trg_exam_reg_bi_eligibility also validates enrolment and attendance, not fee dues. Treat the documented fee gate as unverified until the procedure/trigger is corrected and exercised against MariaDB/MySQL.'
    )),

    B('h1', text='12. Testing and Evidence Status'),
    B('table', headers=['Evidence item', 'Status for this report'], rows=[
        ['Repository/source review', 'Performed: academic UI pages, route file, service, table dictionary, seed fragments, procedure and trigger definitions reviewed.'],
        ['React UI browser rendering', 'Performed with Playwright at 1600×1050 viewport for dashboard and five Group 2 pages; actual React components rendered without a captured page exception.'],
        ['API/backend request', 'Not performed: backend API was unavailable; requests were intercepted in the browser.'],
        ['MySQL/MariaDB integration', 'Not performed: database binaries/service and port 3306 were unavailable.'],
        ['CRUD, procedure or trigger execution', 'Not performed in a live database; static code inspection only.'],
        ['Project test suites', 'Not freshly run during report preparation. Repository summaries are recorded below with their discrepancies.'],
    ]),
    B('h2', text='12.1 Existing repository test summaries (not this run)'),
    B('table', headers=['Source document', 'Reported result', 'Qualification'], rows=[
        ['README.md', '10/10 database consistency checks and 158 API smoke checks reported as passing.', 'README also notes three MariaDB collation-related SQL rule statements remain under investigation.'],
        ['PROJECT_REPORT.md §31', '68 database rule tests, 158 API smoke tests and 40 workflow tests reported as passing.', 'The database-test count differs from the README summary; this report did not rerun either suite.'],
    ]),
    B('p', text=(
        'For a defense or release, capture a fresh test run from the current checkout and target database, keep stdout logs, and reconcile the conflicting database-test counts before presenting them as one result.'
    )),

    B('h1', text='13. Limitations and Recommendations'),
    B('numbered', items=[
        'Complete the generated Module 2 ER diagram: add batches, faculty_subjects and academic_calendar to scripts/generate_er_diagrams.py, update Mermaid source, and regenerate docs/er/. The present report calls out the omission rather than disguising it.',
        'Reconcile the exam-fee gate: either add a server/database-backed dues check to the registration path and test it, or correct documentation and UI claims that imply the gate is enforced.',
        'Run database rule, API smoke and workflow suites against the supported MariaDB/MySQL versions; align README and PROJECT_REPORT test counts.',
        'Add explicit positive/negative tests for duplicate subject/program keys, bulk-enrolment rollback, at-capacity offerings, student/offering mismatch and timetable same-start conflicts.',
        'If true interval-overlap prevention is required, extend the timetable trigger/constraints to compare start/end intervals, not only identical start_time values.',
        'Regenerate browser captures after a DB-backed deployment and replace demo metrics with a clearly identified approved dataset; preserve the real official logo asset and canonical E&TC display name.',
    ]),
    B('p', text=(
        'Other project-level checks, including the database dictionary counts and integrations outside Group 2, should be independently verified before a production or institutional deployment. This report does not claim production readiness.'
    )),

    B('h1', text='14. Conclusion'),
    B('p', text=(
        'Group 2 provides the academic backbone that connects curriculum definition to actual delivery. Its relational model separates masters, curriculum mapping, offerings, enrolments and timetable slots, while database-side guards protect the most consequential capacity and scheduling rules. '
        'The React interface presents these tasks within the shared ERP shell and exposes the academic calendar to the rest of the institution.'
    )),
    B('p', text=(
        'The repository contains substantial DBMS artifacts and working UI routes, but evidence must be stated precisely: screenshots here are React UI rendered with illustrative synthetic responses; stored routines and constraints were inspected but not executed in this environment. The diagram inventory and cross-module fee-gate discrepancy are explicitly documented so they can be resolved before a viva or final submission.'
    )),

    B('h1', text='Appendix A. Recommended Demonstration Test Matrix'),
    B('table', headers=['Case', 'Action', 'Expected outcome', 'Execution status'], rows=[
        ['A1', 'Create two subjects with the same subject_code.', 'Second insert is rejected by the unique constraint and API surfaces a useful error.', 'Recommended; not run here.'],
        ['A2', 'Enroll one student twice in one offering.', 'Second row is refused by uq_enrollments_student_offering.', 'Recommended; not run here.'],
        ['A3', 'Bulk-enrol a cohort into a full offering.', 'Procedure rolls back and returns a capacity error; counters remain accurate.', 'Recommended; not run here.'],
        ['A4', 'Insert a slot with end_time earlier than start_time.', 'trg_timetable_bi_conflict raises a readable error.', 'Recommended; not run here.'],
        ['A5', 'Insert a duplicate faculty/room/section day + start slot.', 'Trigger and unique key refuse the duplicate.', 'Recommended; not run here.'],
        ['A6', 'Attempt an ADMIN-only write as STUDENT.', 'API returns 403/unauthorized without changing database state.', 'Recommended; not run here.'],
        ['A7', 'Mark attendance for a student not enrolled in the session offering.', 'Attendance trigger rejects the row.', 'Recommended; not run here.'],
        ['A8', 'Register for an exam with low attendance and with outstanding dues.', 'Attendance case should be rejected; dues case must be separately fixed/verified before claiming a fee gate.', 'Recommended; not run here.'],
        ['A9', 'Add a calendar event with a non-existent academic year or semester.', 'Foreign-key validation rejects the invalid association.', 'Recommended; not run here.'],
    ]),

    B('pagebreak'),
    B('h1', text='Appendix B. Repository Evidence and References'),
    B('bullets', items=[
        'README.md — project scope, stack, seed-count summaries, feature tour, API conventions, setup and test notes.',
        'PROJECT_REPORT.md — overall project report; especially §§7–9 (stack/architecture), §21 (Group 2), §31–32 (testing/setup).',
        'docs/MODULE_INTEGRATION.md — module reference, shared masters, cross-module flows and test/consistency notes.',
        'docs/DATA_DICTIONARY.md — schema inventory, Group 2 table definitions, indexes and foreign-key relationships.',
        'docs/ER_DIAGRAMS.md; docs/er/er_module2_academic_course.svg; docs/er/er_module2_academic_course.mmd — generated ER artifacts and recorded entity/relationship counts.',
        'database/02_tables.sql; database/03_constraints.sql — table definitions, foreign keys, indexes, unique keys and checks.',
        'database/04_seed.sql — academic year/semester seed, programs, subjects and calendar event examples. It is repository seed data, not a current database query.',
        'database/05_procedures.sql — sp_bulk_enroll_students and exam-registration implementation.',
        'database/06_triggers.sql — enrolment counters/guards, timetable conflict guard and exam eligibility trigger.',
        'backend/src/routes/academic.routes.ts; backend/src/controllers/academic.controller.ts; backend/src/services/academic.service.ts; backend/src/repositories/academic.repository.ts — Group 2 API and business/data-access path.',
        'frontend/src/pages/academics/AcademicsMastersPage.tsx, SubjectsPage.tsx, OfferingsPage.tsx, TimetablePage.tsx, CalendarPage.tsx — rendered module pages.',
        'frontend/src/pages/DashboardPage.tsx; frontend/src/components/dashboard/; frontend/public/brand/pict-logo-256.png — shared dashboard and available official PICT logo.',
        'docs/reports/group2/capture_screenshots.cjs — browser-capture script with synthetic API fixtures used for the figures; docs/reports/group2/screenshots/ — standalone screenshot files.',
    ]),
    B('h2', text='Related deliverable'),
    B('p', text=(
        'The separate Group 2 Academic & Course study guide is stored under docs/viva_guides/. This formal project report is a distinct deliverable and does not replace the existing study guide.'
    )),
]


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=70, start=90, bottom=70, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        tc_pr.append(tc_mar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tc_mar.find(qn(f'w:{m}'))
        if node is None:
            node = OxmlElement(f'w:{m}')
            tc_mar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement('w:tblHeader')
    tbl_header.set(qn('w:val'), 'true')
    tr_pr.append(tbl_header)


def add_docx_page_field(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement('w:fldChar')
    fld_char1.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = ' PAGE '
    fld_char2 = OxmlElement('w:fldChar')
    fld_char2.set(qn('w:fldCharType'), 'end')
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)


def add_docx_table(doc, headers, rows, small=False):
    ncols = len(headers)
    table = doc.add_table(rows=1, cols=ncols)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for j, text in enumerate(headers):
        cell = hdr.cells[j]
        cell.text = str(text)
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in cell.paragraphs:
            p.paragraph_format.space_after = Pt(0)
            for r in p.runs:
                r.bold = True
                r.font.name = 'Aptos'
                r.font.size = Pt(8 if small else 8.8)
                r.font.color.rgb = RGBColor(255,255,255)
    for i, row in enumerate(rows):
        cells = table.add_row().cells
        for j, value in enumerate(row):
            cells[j].text = str(value)
            cells[j].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            set_cell_margins(cells[j])
            if i % 2 == 1:
                set_cell_shading(cells[j], 'F4F7FB')
            for p in cells[j].paragraphs:
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.line_spacing = 1.04
                for r in p.runs:
                    r.font.name = 'Aptos'
                    r.font.size = Pt(7.8 if small else 8.4)
                    r.font.color.rgb = RGBColor(36,52,71)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)
    after.paragraph_format.space_before = Pt(0)
    return table


def add_docx_code(doc, text):
    table = doc.add_table(rows=1, cols=1)
    table.style = 'Table Grid'
    cell = table.cell(0,0)
    set_cell_shading(cell, 'F4F7FB')
    set_cell_margins(cell, 100, 130, 100, 130)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    r.font.name = 'Consolas'
    r.font.size = Pt(8.1)
    r.font.color.rgb = RGBColor(36,52,71)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_docx_image(doc, path, caption, width=6.72, page=False):
    if page:
        doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    p.add_run().add_picture(str(path), width=Inches(width))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(0)
    cap.paragraph_format.space_after = Pt(8)
    r = cap.add_run(caption)
    r.italic = True
    r.font.size = Pt(8.3)
    r.font.color.rgb = RGBColor(91,107,122)


def build_docx():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    section.header_distance = Inches(0.32)
    section.footer_distance = Inches(0.30)
    section.different_first_page_header_footer = True

    styles = doc.styles
    normal = styles['Normal']
    normal.font.name = 'Aptos'
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor(36,52,71)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.12

    for name, size, color in [('Title',30,NAVY),('Subtitle',15,TEAL),('Heading 1',19,NAVY),('Heading 2',13,BLUE)]:
        st = styles[name]
        st.font.name = 'Aptos Display' if name in ('Title','Heading 1','Heading 2') else 'Aptos'
        st.font.size = Pt(size)
        st.font.bold = name != 'Subtitle'
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(10 if name == 'Heading 1' else 7)
        st.paragraph_format.space_after = Pt(5)
        if name.startswith('Heading'):
            st.paragraph_format.keep_with_next = True
    if 'Code Block' not in styles:
        code_style = styles.add_style('Code Block', WD_STYLE_TYPE.PARAGRAPH)
        code_style.font.name = 'Consolas'
        code_style.font.size = Pt(8)
    
    # Running header and footer on all non-cover pages.
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.paragraph_format.space_after = Pt(0)
    hrun = header.add_run('PICT · DBMS MAJOR PROJECT  |  GROUP 2 — ACADEMIC & COURSE')
    hrun.font.name = 'Aptos'
    hrun.font.size = Pt(8)
    hrun.font.color.rgb = RGBColor(91,107,122)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_before = Pt(0)
    fr = footer.add_run('Evidence-based report · 17 October 2026  |  Page ')
    fr.font.size = Pt(8)
    fr.font.color.rgb = RGBColor(91,107,122)
    add_docx_page_field(footer)

    # Cover.
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(24)
    p.paragraph_format.space_after = Pt(13)
    p.add_run().add_picture(str(LOGO), width=Inches(1.25))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run('SCTR’S PUNE INSTITUTE OF COMPUTER TECHNOLOGY')
    r.font.name = 'Aptos Display'; r.font.size = Pt(14); r.bold = True; r.font.color.rgb = RGBColor.from_string(NAVY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run('PUNE, MAHARASHTRA')
    r.font.size = Pt(9); r.font.color.rgb = RGBColor.from_string(MUTED)
    p = doc.add_paragraph(style='Title')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    p.add_run('UNIVERSITY HIGHER EDUCATION ERP SYSTEM')
    p = doc.add_paragraph(style='Subtitle')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(3)
    p.add_run('GROUP 2  ·  ACADEMIC & COURSE MODULE')
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(24)
    r = p.add_run('PROJECT STUDY AND IMPLEMENTATION REPORT')
    r.font.size = Pt(12); r.bold = True; r.font.color.rgb = RGBColor.from_string(BLUE)
    cover = doc.add_table(rows=0, cols=2)
    cover.style = 'Table Grid'
    cover.alignment = WD_TABLE_ALIGNMENT.CENTER
    cover.autofit = False
    cover.columns[0].width = Inches(1.55)
    cover.columns[1].width = Inches(4.85)
    for key,val in [
        ('Course','Database Management Systems — Major Project'),
        ('Batch','M7 BATCH'),
        ('Department','Electronics and Telecommunication Engineering (E&TC)'),
        ('Prepared for','Group 2'),
        ('Prepared by','Group members: ______________________________________________'),
        ('Evidence cut-off','17 October 2026'),
    ]:
        cells = cover.add_row().cells
        cells[0].text = key
        cells[1].text = val
        set_cell_shading(cells[0], PALE)
        for c in cells:
            set_cell_margins(c,130,120,130,120)
            for pp in c.paragraphs:
                pp.paragraph_format.space_after = Pt(0)
                for rr in pp.runs:
                    rr.font.name = 'Aptos'; rr.font.size = Pt(10)
        for rr in cells[0].paragraphs[0].runs:
            rr.bold = True; rr.font.color.rgb = RGBColor.from_string(NAVY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(22)
    r = p.add_run('Repository-grounded · synthetic UI responses explicitly labelled · no live DB verification claimed')
    r.font.size = Pt(8.5); r.italic = True; r.font.color.rgb = RGBColor.from_string(MUTED)
    doc.add_page_break()

    # Contents page.
    p = doc.add_paragraph('Contents', style='Heading 1')
    contents = [
        '1. Executive Summary', '2. Project Context, Problem and Objectives', '3. Architecture and Technology',
        '4. Requirements and User Roles', '5. Group 2 Module Design and Workflows', '6. Data Model and ER Diagram',
        '7. Data Dictionary, Keys and Integrity', '8. Database Procedures, Triggers and Transactions',
        '9. API Surface and Security', '10. User Interface Evidence', '11. Cross-Module Integration',
        '12. Testing and Evidence Status', '13. Limitations and Recommendations', '14. Conclusion',
        'Appendix A. Recommended Demonstration Test Matrix', 'Appendix B. Repository Evidence and References',
    ]
    for item in contents:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.space_after = Pt(5)
        r = p.add_run(item)
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor.from_string(INK)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    r = p.add_run('Reading the evidence labels')
    r.bold = True; r.font.color.rgb = RGBColor.from_string(TEAL)
    p = doc.add_paragraph('Repository facts are attributed to specific source files. Screenshots are authentic React UI captures using illustrative synthetic API responses. Database/API behavior not exercised here is identified as unverified.')
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.right_indent = Inches(0.18)
    p.paragraph_format.space_after = Pt(2)
    doc.add_page_break()

    for block in CONTENT:
        t = block['type']
        if t == 'h1':
            doc.add_paragraph(block['text'], style='Heading 1')
        elif t == 'h2':
            doc.add_paragraph(block['text'], style='Heading 2')
        elif t == 'p':
            p = doc.add_paragraph(block['text'])
            p.paragraph_format.widow_control = True
        elif t == 'bullets' or t == 'numbered':
            sty = 'List Bullet' if t == 'bullets' else 'List Number'
            for item in block['items']:
                p = doc.add_paragraph(item, style=sty)
                p.paragraph_format.space_after = Pt(2.5)
                p.paragraph_format.line_spacing = 1.05
        elif t == 'table':
            add_docx_table(doc, block['headers'], block['rows'], small=len(block['headers']) >= 4)
        elif t == 'callout':
            tb = doc.add_table(rows=1, cols=1)
            tb.style = 'Table Grid'
            cell = tb.cell(0,0)
            set_cell_shading(cell, 'EEF4FA')
            set_cell_margins(cell,130,160,130,160)
            pp = cell.paragraphs[0]
            pp.paragraph_format.space_after = Pt(0)
            rr = pp.add_run(block['text'])
            rr.font.name = 'Aptos'; rr.font.size = Pt(9.2); rr.font.color.rgb = RGBColor.from_string(NAVY)
            doc.add_paragraph().paragraph_format.space_after = Pt(1)
        elif t == 'code':
            add_docx_code(doc, block['text'])
        elif t == 'image':
            add_docx_image(doc, block['path'], block['caption'], block.get('width',6.85), False)
        elif t == 'image_page':
            add_docx_image(doc, block['path'], block['caption'], 6.72, True)
        elif t == 'pagebreak':
            doc.add_page_break()

    # Keep title metadata meaningful.
    doc.core_properties.title = 'PICT ERP — Group 2 Academic & Course Module Report'
    doc.core_properties.subject = 'DBMS major project report, Group 2'
    doc.core_properties.author = 'Group 2 — M7 BATCH, Department of E&TC'
    doc.core_properties.keywords = 'PICT, ERP, Academic, Course, DBMS, Group 2'
    DOCX_OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(DOCX_OUT)


# ----------------------------- PDF builder ----------------------------- #


def register_pdf_fonts():
    regular = Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf')
    bold = Path('/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf')
    mono = Path('/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf')
    if regular.exists() and bold.exists():
        pdfmetrics.registerFont(TTFont('DejaVu', str(regular)))
        pdfmetrics.registerFont(TTFont('DejaVu-Bold', str(bold)))
        pdfmetrics.registerFont(TTFont('DejaVuMono', str(mono)))
        pdfmetrics.registerFontFamily('DejaVu', normal='DejaVu', bold='DejaVu-Bold', italic='DejaVu', boldItalic='DejaVu-Bold')
        return 'DejaVu','DejaVu-Bold','DejaVuMono'
    return 'Helvetica','Helvetica-Bold','Courier'


PDF_REG, PDF_BOLD, PDF_MONO = register_pdf_fonts()


class NumberedDocTemplate(BaseDocTemplate):
    def __init__(self, filename, **kwargs):
        super().__init__(filename, **kwargs)
        left = 17*mm; right = 17*mm; top = 17*mm; bottom = 16*mm
        frame = Frame(left, bottom, A4[0]-left-right, A4[1]-top-bottom, id='normal', leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)
        self.addPageTemplates([PageTemplate(id='main', frames=frame, onPage=self.draw_page)] )

    def draw_page(self, canvas, doc):
        canvas.saveState()
        w,h = A4
        if doc.page > 1:
            canvas.setStrokeColor(colors.HexColor('#D4DDE7'))
            canvas.setLineWidth(0.45)
            canvas.line(17*mm, h-12*mm, w-17*mm, h-12*mm)
            canvas.setFont(PDF_REG, 7.2)
            canvas.setFillColor(colors.HexColor('#5B6B7A'))
            canvas.drawRightString(w-17*mm, h-9.5*mm, 'PICT · DBMS MAJOR PROJECT  |  GROUP 2 — ACADEMIC & COURSE')
            canvas.line(17*mm, 12.5*mm, w-17*mm, 12.5*mm)
            canvas.setFont(PDF_REG, 7.1)
            canvas.drawString(17*mm, 8.5*mm, 'Evidence-based report · 17 October 2026')
            canvas.drawRightString(w-17*mm, 8.5*mm, f'Page {doc.page}')
        canvas.restoreState()


def pdf_styles():
    base = getSampleStyleSheet()
    styles = {}
    styles['body'] = ParagraphStyle('BodyERP', parent=base['BodyText'], fontName=PDF_REG, fontSize=8.85, leading=12.25, textColor=colors.HexColor('#243447'), spaceAfter=5.2, alignment=TA_LEFT, allowWidows=0, allowOrphans=0)
    styles['h1'] = ParagraphStyle('H1ERP', parent=base['Heading1'], fontName=PDF_BOLD, fontSize=15.5, leading=19, textColor=colors.HexColor('#16365C'), spaceBefore=9, spaceAfter=6, keepWithNext=True)
    styles['h2'] = ParagraphStyle('H2ERP', parent=base['Heading2'], fontName=PDF_BOLD, fontSize=10.7, leading=13.5, textColor=colors.HexColor('#2D5F91'), spaceBefore=6.5, spaceAfter=3.5, keepWithNext=True)
    styles['caption'] = ParagraphStyle('CaptionERP', parent=base['BodyText'], fontName=PDF_REG, fontSize=7.7, leading=10, textColor=colors.HexColor('#5B6B7A'), alignment=TA_CENTER, spaceBefore=1, spaceAfter=8)
    styles['bullet'] = ParagraphStyle('BulletERP', parent=styles['body'], leftIndent=13, firstLineIndent=-8, spaceAfter=3, leading=11.5)
    styles['table'] = ParagraphStyle('TableERP', parent=styles['body'], fontName=PDF_REG, fontSize=7.15, leading=9.2, spaceAfter=0, spaceBefore=0, textColor=colors.HexColor('#243447'))
    styles['table_small'] = ParagraphStyle('TableSmallERP', parent=styles['body'], fontName=PDF_REG, fontSize=6.65, leading=8.3, spaceAfter=0, spaceBefore=0, textColor=colors.HexColor('#243447'))
    styles['table_head'] = ParagraphStyle('TableHeadERP', parent=styles['table'], fontName=PDF_BOLD, textColor=colors.white)
    styles['callout'] = ParagraphStyle('CalloutERP', parent=styles['body'], fontName=PDF_REG, fontSize=8.25, leading=11, textColor=colors.HexColor('#16365C'), spaceAfter=0)
    styles['code'] = ParagraphStyle('CodeERP', parent=styles['body'], fontName=PDF_MONO, fontSize=7.2, leading=10, backColor=colors.HexColor('#F4F7FB'), borderColor=colors.HexColor('#D4DDE7'), borderWidth=0.4, borderPadding=7, spaceBefore=2, spaceAfter=7)
    styles['contents'] = ParagraphStyle('ContentsERP', parent=styles['body'], fontSize=9.5, leading=15, leftIndent=5, spaceAfter=2)
    styles['cover_org'] = ParagraphStyle('CoverOrgERP', parent=styles['body'], fontName=PDF_BOLD, fontSize=11.5, leading=15, textColor=colors.HexColor('#16365C'), alignment=TA_CENTER, spaceAfter=3)
    styles['cover_city'] = ParagraphStyle('CoverCityERP', parent=styles['body'], fontSize=8.5, leading=11, textColor=colors.HexColor('#5B6B7A'), alignment=TA_CENTER, spaceAfter=12)
    styles['cover_title'] = ParagraphStyle('CoverTitleERP', parent=styles['body'], fontName=PDF_BOLD, fontSize=22, leading=27, textColor=colors.HexColor('#16365C'), alignment=TA_CENTER, spaceBefore=9, spaceAfter=6)
    styles['cover_sub'] = ParagraphStyle('CoverSubERP', parent=styles['body'], fontName=PDF_BOLD, fontSize=14, leading=18, textColor=colors.HexColor('#0F766E'), alignment=TA_CENTER, spaceAfter=4)
    styles['cover_minor'] = ParagraphStyle('CoverMinorERP', parent=styles['body'], fontName=PDF_BOLD, fontSize=10.2, leading=14, textColor=colors.HexColor('#2D5F91'), alignment=TA_CENTER, spaceAfter=18)
    return styles


PSTYLES = pdf_styles()


def para(text, style='body'):
    return Paragraph(html_escape(str(text)), PSTYLES[style])


def pdf_table(headers, rows, avail_width, small=False):
    tstyle = PSTYLES['table_small'] if small or len(headers) >= 4 else PSTYLES['table']
    data = [[Paragraph(html_escape(str(x)), PSTYLES['table_head']) for x in headers]]
    for row in rows:
        data.append([Paragraph(html_escape(str(x)), tstyle) for x in row])
    n = len(headers)
    if n == 2:
        ratios = [0.27,0.73]
    elif n == 3:
        ratios = [0.22,0.31,0.47]
    elif n == 4:
        ratios = [0.11,0.28,0.48,0.13]
    else:
        ratios = [1/n]*n
    widths = [avail_width*r for r in ratios]
    tbl = Table(data, colWidths=widths, repeatRows=1, hAlign='LEFT', splitByRow=1)
    tbl.setStyle(TableStyle([
        ('BACKGROUND',(0,0),(-1,0),colors.HexColor('#16365C')),
        ('TEXTCOLOR',(0,0),(-1,0),colors.white),
        ('GRID',(0,0),(-1,-1),0.35,colors.HexColor('#D4DDE7')),
        ('VALIGN',(0,0),(-1,-1),'TOP'),
        ('LEFTPADDING',(0,0),(-1,-1),5),
        ('RIGHTPADDING',(0,0),(-1,-1),5),
        ('TOPPADDING',(0,0),(-1,-1),4),
        ('BOTTOMPADDING',(0,0),(-1,-1),4),
        ('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#F4F7FB')]),
    ]))
    return tbl


def build_pdf():
    pdf = NumberedDocTemplate(str(PDF_OUT), pagesize=A4, title='PICT ERP — Group 2 Academic & Course Module Report', author='Group 2 — M7 BATCH, Department of E&TC', subject='DBMS major project report')
    width = A4[0] - 34*mm
    story = []
    # Cover page.
    logo = RLImage(str(LOGO), width=31*mm, height=31*mm)
    story += [Spacer(1,11*mm), logo, Spacer(1,5*mm),
              para('SCTR’S PUNE INSTITUTE OF COMPUTER TECHNOLOGY','cover_org'),
              para('PUNE, MAHARASHTRA','cover_city'),
              Spacer(1,6*mm), para('UNIVERSITY HIGHER EDUCATION ERP SYSTEM','cover_title'),
              para('GROUP 2  ·  ACADEMIC & COURSE MODULE','cover_sub'),
              para('PROJECT STUDY AND IMPLEMENTATION REPORT','cover_minor')]
    meta = [
        [Paragraph('<b>Course</b>',PSTYLES['table']),Paragraph('Database Management Systems — Major Project',PSTYLES['table'])],
        [Paragraph('<b>Batch</b>',PSTYLES['table']),Paragraph('M7 BATCH',PSTYLES['table'])],
        [Paragraph('<b>Department</b>',PSTYLES['table']),Paragraph('Electronics and Telecommunication Engineering (E&amp;TC)',PSTYLES['table'])],
        [Paragraph('<b>Prepared for</b>',PSTYLES['table']),Paragraph('Group 2',PSTYLES['table'])],
        [Paragraph('<b>Prepared by</b>',PSTYLES['table']),Paragraph('Group members: ______________________________________________',PSTYLES['table'])],
        [Paragraph('<b>Evidence cut-off</b>',PSTYLES['table']),Paragraph('17 October 2026',PSTYLES['table'])],
    ]
    mt = Table(meta, colWidths=[38*mm,width-38*mm], hAlign='CENTER')
    mt.setStyle(TableStyle([
        ('BACKGROUND',(0,0),(0,-1),colors.HexColor('#EEF4FA')),
        ('GRID',(0,0),(-1,-1),0.4,colors.HexColor('#D4DDE7')),
        ('VALIGN',(0,0),(-1,-1),'MIDDLE'),
        ('LEFTPADDING',(0,0),(-1,-1),8),('RIGHTPADDING',(0,0),(-1,-1),8),
        ('TOPPADDING',(0,0),(-1,-1),8),('BOTTOMPADDING',(0,0),(-1,-1),8),
    ]))
    story += [mt, Spacer(1,12*mm), Paragraph('Repository-grounded · synthetic UI responses explicitly labelled · no live DB verification claimed',ParagraphStyle('coverNote',fontName=PDF_REG,fontSize=7.6,leading=10,textColor=colors.HexColor('#5B6B7A'),alignment=TA_CENTER)), PageBreak()]
    # Contents page.
    story.append(Paragraph('Contents', PSTYLES['h1']))
    for item in [
        '1. Executive Summary','2. Project Context, Problem and Objectives','3. Architecture and Technology',
        '4. Requirements and User Roles','5. Group 2 Module Design and Workflows','6. Data Model and ER Diagram',
        '7. Data Dictionary, Keys and Integrity','8. Database Procedures, Triggers and Transactions',
        '9. API Surface and Security','10. User Interface Evidence','11. Cross-Module Integration',
        '12. Testing and Evidence Status','13. Limitations and Recommendations','14. Conclusion',
        'Appendix A. Recommended Demonstration Test Matrix','Appendix B. Repository Evidence and References']:
        story.append(Paragraph('•&nbsp;&nbsp;'+html_escape(item), PSTYLES['contents']))
    story += [Spacer(1,5*mm), Paragraph('<b>Reading the evidence labels</b>',PSTYLES['h2']),
              para('Repository facts are attributed to specific source files. Screenshots are authentic React UI captures using illustrative synthetic API responses. Database/API behavior not exercised here is identified as unverified.'), PageBreak()]

    for block in CONTENT:
        t = block['type']
        if t == 'h1':
            story.append(Paragraph(html_escape(block['text']), PSTYLES['h1']))
        elif t == 'h2':
            story.append(Paragraph(html_escape(block['text']), PSTYLES['h2']))
        elif t == 'p':
            story.append(para(block['text']))
        elif t == 'bullets' or t == 'numbered':
            for i,item in enumerate(block['items'],1):
                bullet = '•' if t == 'bullets' else f'{i}.'
                story.append(Paragraph(f'{bullet}&nbsp;&nbsp;{html_escape(item)}',PSTYLES['bullet']))
        elif t == 'table':
            story.append(pdf_table(block['headers'],block['rows'],width,small=len(block['headers'])>=4))
            story.append(Spacer(1,4))
        elif t == 'callout':
            cp = Paragraph(html_escape(block['text']), PSTYLES['callout'])
            ct = Table([[cp]],colWidths=[width],hAlign='LEFT')
            ct.setStyle(TableStyle([
                ('BACKGROUND',(0,0),(-1,-1),colors.HexColor('#EEF4FA')),
                ('BOX',(0,0),(-1,-1),0.6,colors.HexColor('#B7C9DA')),
                ('LEFTPADDING',(0,0),(-1,-1),9),('RIGHTPADDING',(0,0),(-1,-1),9),
                ('TOPPADDING',(0,0),(-1,-1),7),('BOTTOMPADDING',(0,0),(-1,-1),7),
            ]))
            story += [ct,Spacer(1,5)]
        elif t == 'code':
            code_text = html_escape(block['text']).replace('\n','<br/>')
            story += [Paragraph(code_text,PSTYLES['code']),Spacer(1,3)]
        elif t == 'image':
            im = RLImage(str(block['path']))
            im._restrictSize(width, 7.4*72)
            story += [im, Paragraph(html_escape(block['caption']),PSTYLES['caption'])]
        elif t == 'image_page':
            story.append(PageBreak())
            im = RLImage(str(block['path']))
            im._restrictSize(width, 7.55*72)
            story += [im, Paragraph(html_escape(block['caption']),PSTYLES['caption'])]
        elif t == 'pagebreak':
            story.append(PageBreak())
    PDF_OUT.parent.mkdir(parents=True,exist_ok=True)
    pdf.build(story)


if __name__ == '__main__':
    missing = [p for p in [LOGO,ER_IMAGE,*[SCREEN_DIR/f'{i:02d}_{name}.png' for i,name in [(1,'Admin_Dashboard_UI'),(2,'Academic_Masters_UI'),(3,'Subjects_UI'),(4,'Course_Offerings_UI'),(5,'Timetable_UI'),(6,'Academic_Calendar_UI')]]] if not p.exists()]
    if missing:
        raise SystemExit('Missing report assets: ' + ', '.join(str(p) for p in missing))
    build_docx()
    build_pdf()
    print(DOCX_OUT)
    print(PDF_OUT)
